import openpyxl
path = r'D:\长征文化\009 变更记录\长征文化项目EC变更跟踪表.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb.active

# Row 13 (EC #11): 点击漫游没反应 - update analysis and fix
ws['F13'] = '前端/JS'
ws['G13'] = '两个根因：1) showSpeedSelector() 显示 play-controls 时未设置定位坐标（top/left/right），play-controls 只有 position:fixed 未指定位置，出现在页面底部 y=1020（屏幕底边）。2) 代码随后将 speed-selector 定位在 play-controls 下方 y=1108（屏幕外）。用户完全看不到面板，感觉"点了没反应"。'
ws['H13'] = '在 showSpeedSelector() 中给 play-controls 添加定位（和 speed-selector 一样靠按钮定位），再调整 speed-selector 到其下方。已在 DEV 部署并浏览器验证：面板可见、位置正确、gap=4px、点击外部关闭正常、选择速度开始漫游正常。'
ws['I13'] = '已修复'

wb.save(path)
print("EC table updated: #11 -> 已修复")

# Cleanup temp files
import os
for f in os.listdir(r'D:\长征文化'):
    if f.startswith('_') and (f.endswith('.py') or f.endswith('.js') or f.endswith('.css') or f.endswith('.html')) and f not in ['_fix_position.py', '_controls_check.js']:
        os.remove(os.path.join(r'D:\长征文化', f))
        print(f'removed {f}')

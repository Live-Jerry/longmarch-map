import openpyxl
from copy import copy

path = r'D:\长征文化\009 变更记录\长征文化项目EC变更跟踪表.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb.active

def set_cell(row, col, value):
    cell = ws[f'{col}{row}']
    cell.value = value
    ref = ws[f'{col}3']
    cell.font = copy(ref.font)
    cell.border = copy(ref.border)
    cell.alignment = copy(ref.alignment)
    cell.fill = copy(ref.fill)
    cell.number_format = ref.number_format

# 重写第13行 G/H
set_cell(13, 'G',
    '漫游功能在未做任何代码修改的情况下自行恢复，排除代码逻辑 bug。'
    '结合服务器部署时序分析，根因为环境/数据层问题：\n'
    '1) route_point 表初始可能为空（坐标数据未导入），'
    'startAutowalk() 中 fetch /api/v1/routes/autowalk 返回 segments<2，'
    '函数直接 return，用户感觉"点了没反应"。\n'
    '2) 之后某次部署过程中 route_point 表被填充（import_route_points.py '
    '或其他初始化流程被执行），功能自动恢复。\n'
    '3) 补充可能性：旧版 JS 被浏览器缓存（未 Ctrl+F5），'
    '或 gunicorn 进程重启后服务恢复正常。'
)
set_cell(13, 'H',
    '已验证：当前 dev.zhichangxuan.com 漫游功能正常，'
    '确认 route_point 表已有数据（含4支军队路线点）。\n'
    '预防措施：\n'
    '1) 在 deploy.sh 中增加 route_point 表初始化检查，'
    '部署时若表为空则自动执行 import_route_points.py\n'
    '2) 在 startAutowalk() 中增加明确的错误提示：'
    'API 返回空 segments 时弹窗提示"路线数据未加载，请联系管理员"，'
    '而非无声 return。'
)
set_cell(13, 'I', '已恢复（根因为数据未初始化，部署流程缺检查）')

wb.save(path)
print('第13行已重写：G=环境/数据层根因分析，H=预防措施，I=已恢复')

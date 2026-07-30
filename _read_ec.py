import openpyxl
wb = openpyxl.load_workbook(r'D:\长征文化\009 变更记录\长征文化项目EC变更跟踪表.xlsx')
ws = wb.active
for r in range(1, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    d = ws.cell(row=r, column=4).value
    e = ws.cell(row=r, column=5).value
    i = ws.cell(row=r, column=9).value
    if a or d:
        an = str(a) if a else ''
        status = str(i) if i else ''
        dn = str(d)[:80] if d else ''
        en = str(e)[:100] if e else ''
        print(f'#{an} ({status})')
        print(f'  D: {dn}')
        if en:
            print(f'  E: {en}')
        print()

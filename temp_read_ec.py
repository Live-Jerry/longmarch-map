import openpyxl
wb = openpyxl.load_workbook(
    r'D:\长征文化\009 变更记录\长征文化项目EC变更跟踪表.xlsx'
)
ws = wb.active
for row in range(2, ws.max_row + 1):
    seq = ws.cell(row=row, column=1).value
    typ = ws.cell(row=row, column=2).value
    desc = ws.cell(row=row, column=4).value
    res = ws.cell(row=row, column=5).value
    c6 = ws.cell(row=row, column=6).value
    c7 = ws.cell(row=row, column=7).value
    c8 = ws.cell(row=row, column=8).value
    if seq and str(seq).strip().isdigit():
        empty = '<<<' if (not c6 or not c7 or not c8) else ''
        desc_s = (desc[:60] + '..') if desc else '(空)'
        res_s = (res[:60] + '..') if res else '(空)'
        print(f'#{seq} {typ}')
        print(f'  问题: {desc_s}')
        print(f'  预期: {res_s}')
        print(f'  C6={bool(c6)} C7={bool(c7)} C8={bool(c8)} {empty}')
        print()

import openpyxl
path = r'D:\长征文化\009 变更记录\长征文化项目EC变更跟踪表.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb.active

# Row 13 (EC #11): 点击漫游没反应
ws['F13'] = '前端/JS'
ws['G13'] = '可能原因：1) JS 运行时错误阻止 initControls 执行；2) CSS 层叠导致按钮不可见/不可点击；3) 浏览器缓存旧版 JS（nginx no-cache 已设置但用户可能未 Ctrl+F5）；4) controls.js 引用不存在的函数？需检查 initControls 中所有 getElementById 的 DOM 元素是否存在'
ws['H13'] = '需用户在浏览器按 F12 查看 Console 报错。如为缓存问题：Ctrl+F5 强制刷新。如为 JS 错误：在 DEV 服务器上打开 dev.zhichangxuan.com 查看具体报错'
ws['I13'] = '待分析'

# Row 14 (EC #12): 管理页面侧边栏问题
ws['F14'] = '前端/模板'
ws['G14'] = 'nodes/sparks/users/index 等 admin 页面的侧边栏：1) 缺少"留言管理"链接，或链接缩进不正确；2) messages.html 原用 extends/index.html 继承，但 index.html 无 block 定义导致继承损坏'
ws['H14'] = '1) 所有 admin 侧边栏已添加留言管理链接，缩进已修复。2) messages.html 已重写为独立页面（不继承 index.html），直接显示留言。3) 已通过 curl 验证各 admin 页面 HTTP 200'
ws['I14'] = '已修复'

# Update status for #2-#7 from "测试中" to "已修复"
for row_num in range(4, 10):
    ws.cell(row=row_num, column=9).value = '已修复'

wb.save(path)
print("EC table updated")
print("Rows 4-9 (EC#2-#7): -> 已修复")
print("Row 13 (EC#11): -> 待分析")
print("Row 14 (EC#12): -> 已修复")
